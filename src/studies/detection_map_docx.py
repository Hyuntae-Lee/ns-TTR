"""Word report for the detection-limit study (reads cases.csv / summary.json and the figures).

Run from the src directory after detection_map.py and detection_map_report.py:
    .venv\\Scripts\\python.exe studies\\detection_map_docx.py
"""
from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SRC / "studies"))
from detection_map import CRITERIA_DEFAULT, N_SIGMA, min_detectable_half_width  # noqa: E402

OUT = SRC / "results" / "detection_map"
META = json.loads((OUT / "summary.json").read_text(encoding="utf-8"))
FONT = "Malgun Gothic"


def load():
    rows = []
    with open(OUT / "cases.csv", encoding="utf-8-sig") as fh:
        for r in csv.DictReader(fh):
            for k, v in r.items():
                if k in ("group", "preset", "limiting"):
                    continue
                r[k] = (v == "True") if v in ("True", "False") else float(v)
            rows.append(r)
    return rows


ROWS = load()
DEPTHS, WIDTHS = META["depths"], META["half_widths_main"]
GRID = [r for r in ROWS if r["group"] in ("main", "boundary")]
MAIN = {(r["depth_um"], r["half_width_um"]): r for r in ROWS if r["group"] == "main"}
C_TR = abs(META["c_tr"])

OFF = []
if (OUT / "offaxis_cases.csv").exists():
    with open(OUT / "offaxis_cases.csv", encoding="utf-8-sig") as fh:
        for r in csv.DictReader(fh):
            for k, v in r.items():
                if k in ("set", "beam", "limiting"):
                    continue
                r[k] = (v == "True") if v in ("True", "False") else float(v)
            OFF.append(r)


def off(name, offset):
    return next(r for r in OFF if r["set"] == name and r["offset_um"] == offset)


def offset_limit(name):
    """Offset at which the ratio falls to C_min (log-linear interpolation between computed offsets)."""
    import math
    pts = sorted([r for r in OFF if r["set"] == name], key=lambda r: r["offset_um"])
    c = CRITERIA_DEFAULT["c_min"]
    for a, b in zip(pts[:-1], pts[1:]):
        if a["ratio"] >= c > b["ratio"]:
            f = (math.log(a["ratio"]) - math.log(c)) / (math.log(a["ratio"]) - math.log(b["ratio"]))
            return a["offset_um"] + f * (b["offset_um"] - a["offset_um"])
    return None

# ----------------------------------------------------------------------------- document helpers
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(2.2)
TEXT_W = 21.0 - 4.4


def set_font(style, size, bold=False, color=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor(*color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), FONT)


set_font(doc.styles["Normal"], 10.5)
doc.styles["Normal"].paragraph_format.space_after = Pt(6)
doc.styles["Normal"].paragraph_format.line_spacing = 1.25
set_font(doc.styles["Title"], 20, bold=True, color=(0x1F, 0x3A, 0x5F))
set_font(doc.styles["Heading 1"], 15, bold=True, color=(0x1F, 0x3A, 0x5F))
set_font(doc.styles["Heading 2"], 12.5, bold=True, color=(0x1F, 0x3A, 0x5F))
set_font(doc.styles["List Bullet"], 10.5)
doc.styles["Heading 1"].paragraph_format.space_before = Pt(16)
doc.styles["Heading 2"].paragraph_format.space_before = Pt(10)


def p(text="", bold=False, size=None, color=None, align=None, italic=False):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold, run.italic = bold, italic
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    if align:
        para.alignment = align
    return para


def bullets(items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def table(header, rows, widths_cm=None, font_size=9.5, align_right_from=1, fills=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(font_size)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade(c, "E6ECF4")
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = ""
            r = cells[j].paragraphs[0].add_run(str(v))
            r.font.size = Pt(font_size)
            cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT if j >= align_right_from else WD_ALIGN_PARAGRAPH.LEFT
            if fills and fills[i][j]:
                shade(cells[j], fills[i][j])
    if widths_cm:
        for row in t.rows:
            for j, w in enumerate(widths_cm):
                row.cells[j].width = Cm(w)
    for row in t.rows:
        for c in row.cells:
            c.paragraphs[0].paragraph_format.space_after = Pt(1)
            c.paragraphs[0].paragraph_format.line_spacing = 1.0
    doc.add_paragraph()
    return t


def figure(fname, caption):
    doc.add_picture(str(OUT / fname), width=Cm(TEXT_W))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = p(caption, size=9.5, color=(0x55, 0x55, 0x55), align=WD_ALIGN_PARAGRAPH.CENTER)
    cap.paragraph_format.space_after = Pt(10)


def g3(x):
    return "%.3g" % x


def pct(x):
    v = abs(x) * 100
    return ("%.0f" % v) if v >= 100 else ("%.1f" % v) if v >= 1 else ("%.2f" % v) if v >= 0.01 else "<0.01"


def um(b):
    if b["bound"] == "none":
        return "검출 불가"
    if b["min_half_width_um"] > META["rod_radius_um"]:
        return "단면 전체 차단 필요"
    return ("≤ " if b["bound"] == "<=" else "") + "%.1f" % b["min_half_width_um"]


def tsec(t):
    return "%.3g ns" % (t * 1e9) if t < 1e-6 else "%.3g µs" % (t * 1e6) if t < 1e-3 else "%.3g ms" % (t * 1e3)


BOUND = {name: {d: min_detectable_half_width([r for r in GRID if r["depth_um"] == d], crit) for d in DEPTHS}
         for name, crit in META["criteria_variants"].items()}
NAMES = list(META["criteria_variants"])
B0 = BOUND[NAMES[0]]

# ============================================================================= title
doc.add_paragraph("ns-TTR void 검출 한계 맵 시뮬레이션 결과 보고", style="Title")
p("TGV(Cu) 내부 공기 void · 깊이 × 가로 반폭 검출 가능 영역 · 2026-10-02", color=(0x55, 0x55, 0x55))
p("이 보고서의 모든 수치는 ns-TTR 시뮬레이터(ttr_sim, GUI 와 같은 계산 코어)로 계산한 결과이며, 측정값이 아닙니다. "
  "판정 기준(σ, C_min, 허용 온도 상승)은 요청서의 가정값입니다.", size=9.5, color=(0x55, 0x55, 0x55))

# ============================================================================= 1 summary
doc.add_heading("1. 요약", level=1)
d_small = [d for d in DEPTHS if B0[d]["min_half_width_um"] is not None and B0[d]["min_half_width_um"] <= 10]
bullets([
    f"기본 기준(σ = 1×10⁻⁶, 3σ, C_min = 2 %, 허용 표면 온도 상승 10 K)에서 검출 가능한 최소 반폭은 깊이 5 µm 에서 "
    f"{um(B0[5])} µm, 20 µm 에서 {um(B0[20])} µm, 100 µm 에서 {um(B0[100])} µm, 400 µm 에서 {um(B0[400])} µm 입니다 "
    f"(void 두께 10 µm, 구리 기둥 반경 40 µm).",
    "깊이 100 µm 이상에서는 void 가 기둥 반경의 절반(반폭 약 20 µm) 이상을 막아야 검출됩니다. "
    "깊이 400 µm 에서는 반폭 약 30 µm 이상, 즉 단면적의 절반 이상을 막아야 합니다.",
    "기본 기준에서는 모든 깊이에서 '비율 조건(C_min)' 이 판정을 결정합니다. 잡음 조건은 여유가 있어, "
    "σ 를 1×10⁻⁷ 로 낮추거나 허용 온도를 50 K 로 올려도 검출 한계는 달라지지 않습니다.",
    "반대로 σ 가 1×10⁻⁵ 이거나 허용 온도 상승이 1 K 이면 잡음 조건이 한계가 되어 최소 반폭이 커집니다 "
    f"(깊이 20 µm: {um(B0[20])} → {um(BOUND['σ = 1e-5'][20])} µm, 깊이 100 µm: {um(B0[100])} → {um(BOUND['σ = 1e-5'][100])} µm).",
    "void 가 단면을 완전히 막으면(반폭 ≥ 40 µm) 깊이 400 µm 에서도 비율 87 % 이상으로 뚜렷하게 검출됩니다. "
    "반폭 35 µm 와 40 µm 사이에서 신호가 수 배에서 10배 이상 뛰므로(깊이 400 µm: 0.52 → 6.55 mK), 마지막 우회로가 닫히는지가 결정적입니다.",
    "세로 두께는 영향이 작습니다. 두께를 1 → 40 µm 로 40배 바꿔도 신호 변화는 3배 이내입니다.",
    "void 가 기둥 축에서 벗어나면 얕아도 검출이 어렵습니다. 깊이 5 µm · 반폭 10 µm void 는 축 위에서 비율 %s %% 이지만 "
    "중심에서 25 µm 벗어나면 %s %% 로 떨어져 검출되지 않습니다. 검출 한계는 중심 거리 약 %.0f µm 입니다(펌프 반경 10 µm, 프로브 5 µm 기준, 4.7절)."
    % (pct(off("depth5", 0)["ratio"]), pct(off("depth5", 25)["ratio"]), offset_limit("depth5")),
])

doc.add_heading("깊이별 검출 가능한 최소 반폭 (기본 기준)", level=2)
table(["깊이 [µm]", "펄스폭", "최소 반폭 [µm]", "차단 면적 비율", "판정을 결정한 조건"],
      [[d, tsec(next(r["tau_p_s"] for r in GRID if r["depth_um"] == d)), um(B0[d]),
        ("%.1f %%" if B0[d]["min_half_width_um"] < 12.6 else "%.0f %%") % (min(B0[d]["min_half_width_um"], 40.0) ** 2 / 40.0 ** 2 * 100)
        + (" 이하" if B0[d]["bound"] == "<=" else ""),
        B0[d]["limiting"] or "— (계산한 최소 반폭 2 µm 도 검출)"] for d in DEPTHS],
      widths_cm=[2.2, 2.6, 3.0, 3.0, 5.8])
p("최소 반폭은 계산한 반폭들 사이를 로그 보간한 값입니다. 차단 면적 비율 = (반폭 / 기둥 반경 40 µm)².", size=9.5, color=(0x55, 0x55, 0x55))

# ============================================================================= 2 stage 0
doc.add_heading("2. 0단계 확인 결과", level=1)
doc.add_heading("2.1 펄스 에너지와 스폿 크기", level=2)
bullets([
    "펌프: 532 nm, Gaussian 빔, 1/e² 반경 %g µm, 사각 펄스. 구리 반사율 R = %.2f (입사 에너지의 %.0f %% 흡수)."
    % (META["pump_w_um"], META["reflectivity"], (1 - META["reflectivity"]) * 100),
    "프로브: 632.8 nm, 1/e² 반경 %g µm 로 가중 평균한 표면 온도. (dR/dT)/R = −1.5×10⁻⁴ /K." % META["probe_w_um"],
    "펄스 에너지는 깊이 구간(프리셋)마다 다른 기본값을 씁니다. 아래 표의 값이며, 이번 계산은 모두 이 기본값으로 수행한 뒤 "
    "후처리로 스케일링했습니다.",
])
used = sorted({(r["preset"], r["tau_p_s"], r["energy_J"]) for r in GRID}, key=lambda x: x[1])
rows = []
for preset, tau, E in used:
    base = next(r for r in GRID if r["preset"] == preset)
    s10 = 10.0 / base["dT_base_max_K"]
    rows.append([preset, tsec(tau), "%g nJ" % (E * 1e9), "%.3g mJ/cm²" % (2 * E / (3.141592653589793 * (10e-6) ** 2) / 10),
                 "%.3g K" % base["dT_base_max_K"], "%.3g µJ" % (E * s10 * 1e6), "%.3g W" % (E * s10 / tau)])
table(["깊이 구간 [µm]", "펄스폭", "기본 에너지", "중심 fluence", "baseline 피크 ΔT", "10 K 에 필요한 에너지", "그때 피크 파워"],
      rows, widths_cm=[2.3, 2.0, 2.0, 2.4, 2.4, 2.8, 2.4], font_size=9)
p("오른쪽 두 열은 baseline 피크를 10 K 로 맞추는 데 필요한 입사 펄스 에너지와 피크 파워(= 에너지 ÷ 펄스폭)입니다. "
  "깊은 구간일수록 필요한 에너지가 커지므로, 실제 장비 출력으로 가능한지 확인이 필요합니다.", size=9.5, color=(0x55, 0x55, 0x55))

doc.add_heading("2.2 선형성", level=2)
p("열 모델은 선형입니다. 물성값(k, ρc, R)이 온도에 의존하지 않습니다. 같은 케이스(깊이 20 µm, 반폭 20 µm, 두께 10 µm)를 "
  "에너지 1배와 3배로 계산한 결과 baseline 피크는 3.000000000000 배, void − baseline 은 2.999999999993 배였고 비율은 같았습니다. "
  "따라서 2단계의 에너지 스케일링은 후처리로 수행했습니다.")

doc.add_heading("2.3 비율의 정의와 부호", level=2)
bullets([
    "비율 = (void − baseline) ÷ (같은 시각의 baseline 온도 상승). 시각은 |void − baseline| 이 최대가 되는 순간입니다. "
    "분모는 baseline 의 최댓값이 아닙니다.",
    "ΔT 기준으로 계산하며, ΔR/R 로 계산해도 열반사 계수가 약분되어 같은 값입니다.",
    "2 µm void 케이스(두께 1 µm, 반폭 4 µm, 69 ns)는 void − baseline = +0.205 K, 그 시각의 baseline = 0.545 K 이므로 "
    "비율은 +37.6 % 입니다. 시뮬레이터에는 이 값이 음수가 되는 경로가 없습니다. 발표자료 5번 슬라이드의 "
    "'−37.6 %', '−5.9 %' 는 부호를 +로 고쳐야 합니다.",
    "이번 스윕의 %d개 케이스에서 void − baseline 과 비율이 음수인 경우는 %d건입니다." % (
        len(ROWS), sum(1 for r in ROWS if r["void_minus_base_K"] < 0 or r["ratio"] < 0)),
])

doc.add_heading("2.4 계산 시간", level=2)
p("한 케이스(baseline + void)에 평균 %.1f 초가 걸려, 전체 %d 케이스를 %.0f 초에 계산했습니다. 거친 격자로 나눠 돌릴 필요가 없어, "
  "오히려 경계를 찾기 위한 반폭을 추가했습니다(3.1절)." % (META["wall_s"] / META["n_cases"], META["n_cases"], META["wall_s"]))

# ============================================================================= 3 method
doc.add_heading("3. 시뮬레이션 조건과 판정 기준", level=1)
doc.add_heading("3.1 스윕 조건", level=2)
bullets([
    "주 격자: 깊이 %s µm × 가로 반폭 %s µm, 두께 %g µm 고정. void 는 축 위의 공기(k = 0.026 W/m·K) 회전타원체." % (
        ", ".join("%g" % d for d in DEPTHS), ", ".join("%g" % w for w in WIDTHS), META["thickness"]),
    "경계를 정밀하게 찾기 위해 반폭 %s µm 를 추가로 계산했습니다(그림의 맵에는 주 격자만 표시)." % ", ".join(
        "%g" % w for w in META["half_widths_extra"]),
    "펄스폭: 각 깊이가 속한 구간(하한 ≤ 깊이 < 상한)의 기존 규칙값. 깊이 2 µm 는 2~5 µm 구간(69 ns), 400 µm 는 400~500 µm 구간(2.76 ms).",
    "보조 스윕(두께): 깊이 20, 100, 400 µm × 반폭 20, 40 µm × 두께 1, 10, 40 µm.",
    "선택 스윕(펄스폭): 깊이 20, 100, 400 µm × 반폭 20, 40 µm 에서 펄스폭을 규칙값의 0.25, 0.5, 1, 2, 4배.",
    "추가 스윕(축에서 벗어난 void): 깊이 5 µm(반폭 10, 두께 2.5 µm)에서 중심 거리 0~30 µm, 깊이 20 µm(반폭 10, 두께 10 µm)에서 0~25 µm. "
    "펌프와 프로브는 축 위에 고정. 3차원 계산이며 한 케이스에 2~20분이 걸립니다.",
])
doc.add_heading("3.2 요청서와 달라진 점 (사전 확인 완료)", level=2)
bullets([
    "반폭 160 µm 대신 30 µm 를 넣었습니다. 구리 기둥 반경이 40 µm 여서 반폭 80 과 160 µm 는 둘 다 '단면을 완전히 막는 void' 가 되어 "
    "거의 같은 결과가 나오기 때문입니다. 반폭 80 µm 열은 '단면 전체 차단' 을 뜻합니다.",
    "신호 피크가 관측 시간창 끝에 걸리는 케이스는 창을 4배씩 늘려 실제 피크를 찾았습니다. 해당 케이스: "
    + ", ".join(sorted({"깊이 %g · 반폭 %g" % (r["depth_um"], r["half_width_um"]) for r in ROWS if r["window_factor"] > 1 and r["group"] == "main"}))
    + " µm (CSV 의 window_factor 열).",
    "ΔT_base,max 는 프로브 가중 평균(시뮬레이터 표의 'baseline 피크') 기준입니다. 빔 중심의 최고 온도는 이보다 %.0f~%.0f %% 높습니다."
    % ((min(r["dT_base_center_max_K"] / r["dT_base_max_K"] for r in GRID) - 1) * 100,
       (max(r["dT_base_center_max_K"] / r["dT_base_max_K"] for r in GRID) - 1) * 100),
])
doc.add_heading("3.3 에너지 스케일링과 판정 기준", level=2)
p("모델이 선형이므로 baseline 피크가 허용 온도 상승이 되도록 에너지를 비례 조정했습니다.")
bullets([
    "s = ΔT_max_allowed ÷ ΔT_base,max",
    "스케일된 (void − baseline) = s × (void − baseline)",
    "스케일된 ΔR/R 차이 = 1.5×10⁻⁴ × 스케일된 (void − baseline)",
])
table(["조건", "기본값", "민감도 분석 값"],
      [["잡음 조건: 스케일된 ΔR/R 차이 ≥ 3σ", "σ = 1×10⁻⁶", "σ = 1×10⁻⁵, 1×10⁻⁷"],
       ["계통 오차 조건: |비율| ≥ C_min", "C_min = 2 %", "1 %, 5 %"],
       ["허용 표면 온도 상승 ΔT_max_allowed", "10 K (미확정, 가정값)", "1 K, 50 K"]],
      widths_cm=[7.2, 4.2, 5.0], align_right_from=9)
p("두 조건을 모두 만족하면 '검출 가능' 으로 판정했습니다. 민감도 분석은 한 번에 한 값만 바꿨습니다. "
  "기준값은 studies/detection_map.py 의 CRITERIA_DEFAULT / CRITERIA_VARIANTS 에서 바꿀 수 있고, 재계산 없이 판정만 다시 할 수 있습니다.")

# ============================================================================= 4 results
doc.add_heading("4. 결과", level=1)
doc.add_heading("4.1 비율 맵", level=2)
figure("fig1_ratio_map.png", "그림 1. 비율 (void − baseline) / baseline [%]. 가로축 반폭, 세로축 깊이 (둘 다 로그 축).")
bullets([
    "같은 반폭에서 깊이가 깊어질수록 비율이 급격히 줄어듭니다. 반폭 20 µm 기준 5 µm 에서 %s %%, 20 µm 에서 %s %%, 100 µm 에서 %s %%, 400 µm 에서 %s %%."
    % tuple(pct(MAIN[(d, 20)]["ratio"]) for d in (5, 20, 100, 400)),
    "반폭 30 µm 까지는 깊이에 따라 완만하게 변하지만, 반폭 40 µm(기둥 벽에 닿음)에서 비율이 한 자릿수 이상 뜁니다. "
    "깊이 100 µm 에서 반폭 35 µm 는 %s %%, 40 µm 는 %s %%." % (
        pct(next(r for r in GRID if r["depth_um"] == 100 and r["half_width_um"] == 35)["ratio"]), pct(MAIN[(100, 40)]["ratio"])),
    "단면을 막는 void 의 비율이 100 % 를 넘는 것은 신호 피크가 baseline 이 이미 많이 식은 뒤에 나타나기 때문입니다 "
    "(분모가 같은 시각의 baseline).",
])
doc.add_heading("4.2 스케일된 ΔR/R 차이 맵", level=2)
figure("fig2_scaled_dRR_map.png", "그림 2. baseline 피크를 10 K 로 맞췄을 때의 |ΔR/R 차이| [×10⁻⁶]. 3σ = 3 이상이면 잡음 조건 만족.")
bullets([
    "허용 온도 상승 10 K 에서는 비율 조건(2 %%)을 만족하는 케이스가 모두 잡음 조건(3×10⁻⁶)도 만족합니다. "
    "깊이 400 µm · 반폭 20 µm 처럼 비율이 0.6 %%인 케이스도 ΔR/R 차이는 %.1f×10⁻⁶ 로 잡음 조건을 넘습니다."
    % (MAIN[(400, 20)]["scaled_dRR_diff"] * 1e6),
    "따라서 기본 기준에서 검출 한계를 정하는 것은 신호의 절대 크기가 아니라, baseline 대비 비율이 2 % 에 못 미친다는 점입니다. "
    "반폭 10 µm 이하의 깊은 void 는 두 조건을 모두 만족하지 못합니다.",
])
doc.add_heading("4.3 검출 가능 영역", level=2)
figure("fig3_detectability_map.png", "그림 3. 기본 기준의 검출 가능(파랑) / 불가(주황, 빗금) 영역과 민감도 기준의 경계선.")
doc.add_heading("4.4 깊이별 검출 가능한 최소 반폭", level=2)
figure("fig4_min_half_width.png", "그림 4. 깊이별 검출 가능한 최소 반폭. 실선: 기본 기준, 점선·파선: 민감도 기준.")
table(["깊이 [µm]"] + [n.replace("기본 (σ=1e-6, C=2 %, 10 K)", "기본 기준") for n in NAMES],
      [[d] + [um(BOUND[n][d]) for n in NAMES] for d in DEPTHS], font_size=9,
      widths_cm=[1.8, 2.2, 2.0, 2.0, 2.2, 2.2, 2.2, 2.2])
p("단위 µm. σ = 1×10⁻⁷ 와 ΔT 허용 50 K 는 기본 기준과 같은 값입니다(잡음 조건에 이미 여유가 있음). "
  "σ = 1×10⁻⁵ 와 ΔT 허용 1 K 는 서로 같은 값입니다(둘 다 잡음 여유를 10배 줄이는 효과).", size=9.5, color=(0x55, 0x55, 0x55))

doc.add_heading("4.5 두께의 영향 (보조 스윕)", level=2)
figure("fig5_thickness.png", "그림 5. 세로 두께 1, 10, 40 µm 에서의 비율(왼쪽)과 스케일된 ΔR/R 차이(오른쪽).")
TH = [r for r in ROWS if r["group"] == "thickness"]
table(["깊이 [µm]", "반폭 [µm]", "두께 [µm]", "void − baseline [mK]", "비율 [%]", "스케일된 ΔR/R 차이", "판정"],
      [["%g" % r["depth_um"], "%g" % r["half_width_um"], "%g" % r["thickness_um"], g3(r["void_minus_base_K"] * 1e3), pct(r["ratio"]),
        "%.2e" % r["scaled_dRR_diff"], "가능" if r["detectable"] else "불가 (%s)" % r["limiting"]] for r in TH],
      font_size=9, widths_cm=[1.8, 1.8, 1.8, 3.2, 2.0, 3.2, 2.8])
bullets([
    "단면을 막는 void(반폭 40 µm)는 두께와 거의 무관합니다. 1 µm 두께의 얇은 틈도 두꺼운 void 와 같은 수준의 신호를 냅니다.",
    "우회로가 남는 void(반폭 20 µm)는 깊이에 따라 방향이 다릅니다. 깊이 100 µm 에서는 두꺼울수록 신호가 커지고(1.7 → 4.4 %), "
    "깊이 20 µm 에서는 오히려 줄어듭니다(17.9 → 12.1 %). 깊이를 void 윗면으로 정의했기 때문에, 두꺼운 타원체일수록 가장 넓은 부분이 더 깊은 곳에 놓이기 때문입니다.",
    "깊이 100 µm · 반폭 20 µm 는 두께 1 µm 에서 비율 1.7 % 로 기준에 못 미치고, 10 µm 이상에서 넘습니다. 검출 경계 근처에서는 두께가 판정을 바꿀 수 있습니다.",
])

doc.add_heading("4.6 펄스폭의 영향 (선택 스윕)", level=2)
figure("fig6_pulse_width.png", "그림 6. 펄스폭을 규칙값의 0.25~4배로 바꿨을 때의 변화 (규칙값 1× 대비 배수). 표면 온도 상승은 10 K 로 고정.")
PU = [r for r in ROWS if r["group"] == "pulse"]
rows = []
for d in META["pulse_sweep"]["depths"]:
    for w in META["pulse_sweep"]["half_widths"]:
        pts = sorted([r for r in PU if r["depth_um"] == d and r["half_width_um"] == w], key=lambda r: r["tau_factor"])
        ref = next(x for x in pts if x["tau_factor"] == 1.0)
        rows.append(["%g" % d, "%g" % w] + ["%.2f / %s" % (x["scaled_dRR_diff"] / ref["scaled_dRR_diff"], pct(x["ratio"])) for x in pts])
table(["깊이 [µm]", "반폭 [µm]"] + ["%g×" % f for f in META["pulse_sweep"]["factors"]], rows, font_size=9,
      widths_cm=[1.8, 1.8, 2.6, 2.6, 2.6, 2.6, 2.6])
p("각 칸: 스케일된 ΔR/R 차이의 1× 대비 배수 / 비율 [%].", size=9.5, color=(0x55, 0x55, 0x55))
bullets([
    "표면 온도 상승을 같은 값(10 K)으로 제한하면, 펄스가 길수록 스케일된 ΔR/R 차이가 커집니다(4배 펄스에서 1.3~3.4배). "
    "같은 표면 온도에서 긴 펄스가 더 많은 에너지를 넣을 수 있기 때문입니다.",
    "반면 우회로가 남는 얕은 void(깊이 20 µm · 반폭 20 µm)는 펄스가 길어지면 비율이 떨어집니다(1× 16.8 % → 4× 10.1 %). "
    "깊은 void(400 µm)는 비율도 함께 올라갑니다.",
    "즉 현재 규칙(구간 하한 기준)은 '비율' 관점에서는 얕은 쪽에 맞는 보수적인 선택이고, '절대 신호' 관점의 최적은 아닙니다. "
    "잡음이 한계인 조건(σ 가 크거나 허용 온도가 낮을 때)에서는 규칙값보다 2~4배 긴 펄스가 유리할 수 있습니다. "
    "다만 긴 펄스는 필요한 펄스 에너지도 커집니다.",
])

doc.add_heading("4.7 축에서 벗어난 void (추가 스윕)", level=2)
p("지금까지의 결과는 모두 void 가 기둥 축 위에 있는 경우입니다. 수동 시뮬레이션에서 '깊이 5 µm 프리셋에서 중심 거리를 25 µm 로 하면 "
  "검출되지 않는다' 는 관찰이 있어, void 를 축에서 옆으로 옮기며 계산했습니다. 펌프(반경 10 µm)와 프로브(반경 5 µm)는 축 위에 그대로 둡니다.")
figure("fig7_offaxis.png", "그림 7. void 중심이 축에서 벗어난 거리에 따른 비율(왼쪽)과 스케일된 ΔR/R 차이(오른쪽). 마커: 펌프를 넓힌 경우.")
table(["조건", "중심 거리 [µm]", "void − baseline [mK]", "비율 [%]", "스케일된 ΔR/R 차이", "판정"],
      [[{"depth5": "깊이 5 µm", "depth20": "깊이 20 µm"}[r["set"]], "%g" % r["offset_um"], g3(r["void_minus_base_K"] * 1e3), pct(r["ratio"]),
        "%.2e" % r["scaled_dRR_diff"], "가능" if r["detectable"] else "불가 (%s)" % r["limiting"]]
       for r in OFF if r["set"] in ("depth5", "depth20")],
      font_size=9, widths_cm=[2.6, 2.6, 3.4, 2.2, 3.4, 2.8])
bullets([
    "깊이 5 µm void 는 중심에서 벗어날수록 신호가 급격히 줄어듭니다. 축 위 %s mK(비율 %s %%)에서 거리 15 µm 는 %s mK(%s %%), 25 µm 는 %s mK(%s %%), "
    "30 µm 는 %s mK(%s %%)입니다. 거리 25 µm 에서 신호는 축 위의 약 1/%d 입니다."
    % (g3(off("depth5", 0)["void_minus_base_K"] * 1e3), pct(off("depth5", 0)["ratio"]),
       g3(off("depth5", 15)["void_minus_base_K"] * 1e3), pct(off("depth5", 15)["ratio"]),
       g3(off("depth5", 25)["void_minus_base_K"] * 1e3), pct(off("depth5", 25)["ratio"]),
       g3(off("depth5", 30)["void_minus_base_K"] * 1e3), pct(off("depth5", 30)["ratio"]),
       round(off("depth5", 0)["void_minus_base_K"] / off("depth5", 25)["void_minus_base_K"])),
    "기본 기준의 검출 한계는 깊이 5 µm 에서 중심 거리 약 %.0f µm, 깊이 20 µm 에서 약 %.0f µm 입니다. void 의 반폭이 10 µm 이므로, "
    "void 의 안쪽 가장자리가 펌프 반경(10 µm) 바깥으로 나가는 지점과 거의 일치합니다."
    % (offset_limit("depth5"), offset_limit("depth20")),
    "이유는 가열 범위입니다. 펌프 빔의 세기는 중심에서 10 µm 떨어지면 13.5 %%, 20 µm 에서는 0.03 %% 로 줄어, 얕은 void 가 빔 밖에 있으면 "
    "void 위쪽이 거의 가열되지 않습니다. 열이 옆으로 퍼져 void 에 닿을 때는 이미 약해져 있고, 그 영향이 다시 축 위의 프로브까지 돌아와야 합니다." % (),
    "깊은 void 는 축 위 신호 자체는 작지만 위치에 덜 민감합니다. 깊이 20 µm 는 축 위 %s %% 에서 거리 15 µm 에 %s %%, 25 µm 에 %s %% 로 완만하게 줄어듭니다. "
    "열이 그 깊이에 도달할 때는 이미 옆으로 넓게 퍼져 있기 때문입니다."
    % (pct(off("depth20", 0)["ratio"]), pct(off("depth20", 15)["ratio"]), pct(off("depth20", 25)["ratio"])),
])
doc.add_heading("펌프를 넓히면 (거리 25 µm 의 깊이 5 µm void)", level=2)
table(["펌프", "프로브 반경", "void 위치", "void − baseline [mK]", "비율 [%]", "스케일된 ΔR/R 차이", "판정"],
      [[lab, pr, "축 위" if o == 0 else "거리 25 µm", g3(off(name, o)["void_minus_base_K"] * 1e3), pct(off(name, o)["ratio"]),
        "%.2e" % off(name, o)["scaled_dRR_diff"], "가능" if off(name, o)["detectable"] else "불가 (%s)" % off(name, o)["limiting"]]
       for name, lab, pr in (("depth5", "Gaussian 10 µm (기본)", "5 µm"), ("flat40_probe5", "Flat-top 40 µm", "5 µm"),
                             ("flat40_probe40", "Flat-top 40 µm", "40 µm"))
       for o in (0, 25)],
      font_size=9, widths_cm=[3.6, 2.0, 2.2, 3.0, 1.8, 2.8, 2.4], align_right_from=3)
bullets([
    "기둥 전체를 고르게 가열하는 flat-top 펌프(반경 40 µm)로 바꾸면 거리 25 µm void 의 비율이 %s → %s %% 로 올라가지만, 프로브가 축 위(반경 5 µm)에 "
    "그대로 있으면 기준 2 %% 에 조금 못 미칩니다." % (pct(off("depth5", 25)["ratio"]), pct(off("flat40_probe5", 25)["ratio"])),
    "프로브도 기둥 전체(반경 40 µm)로 넓히면 %s %% 가 되어 검출됩니다. 대신 축 위 void 의 비율은 %s → %s %% 로 낮아집니다. "
    "넓은 프로브는 위치에 덜 민감한 대신 신호를 평균해 대비가 줄어듭니다."
    % (pct(off("flat40_probe40", 25)["ratio"]), pct(off("flat40_probe5", 0)["ratio"]), pct(off("flat40_probe40", 0)["ratio"])),
    "따라서 축에서 벗어난 얕은 void 를 잡으려면 펌프만 넓혀서는 부족하고, 프로브 범위를 함께 넓히거나 펌프·프로브를 기둥 단면 안에서 옮겨 가며(스캔) 측정해야 합니다.",
    "같은 표면 온도 상승(10 K)을 넓은 빔으로 만들려면 더 큰 펄스 에너지가 필요합니다. 이 비교는 깊이 5 µm 한 케이스이므로, 넓은 빔 조건의 전체 검출 맵은 별도 스윕이 필요합니다.",
])

# ============================================================================= 5 notes
doc.add_heading("5. 예상과 다르거나 주의할 결과", level=1)
bullets([
    "부호가 뒤집히는 케이스는 없었습니다. 모든 케이스에서 void 가 있는 쪽이 더 뜨겁습니다.",
    "반폭 35 → 40 µm 사이의 급변: 깊이 100 µm 에서 비율이 14 % → 173 % 로 뜁니다. 타원체 가장자리가 기둥 벽에 닿아 우회로가 닫히는 "
    "지점이라 물리적으로 타당하지만, 격자 해상도와 void 형상 가정(타원체)에 민감한 구간입니다. 실제 void 가 벽에서 몇 µm 떨어져 있는지에 따라 "
    "결과가 크게 달라집니다.",
    "반폭 40 µm 와 80 µm 의 차이(깊이 400 µm: 87 % vs 94 %)는 타원체 가장자리가 얇아 생기는 미세한 우회로 때문입니다.",
    "단면을 막는 void 는 신호 피크가 매우 늦습니다. 깊이 20 µm · 반폭 40 µm 는 기본 시간창(34.5 µs) 안에서 피크가 나오지 않아 창을 4배로 늘렸습니다. "
    "실험에서도 이런 void 는 긴 관측 시간이 필요합니다.",
    "깊이 2 µm · 반폭 20 µm 이상 케이스의 비율이 100 % 를 넘는 것은 void 가 표면 바로 아래를 넓게 막아, 그 시각의 표면 온도가 baseline 의 2배 이상이 되기 때문입니다.",
])

doc.add_heading("6. 전제와 한계", level=1)
bullets([
    "허용 표면 온도 상승 10 K, 잡음 σ = 1×10⁻⁶, C_min = 2 % 는 가정값입니다. 실제 장비 출력, 시료 손상 한계, 측정 재현성으로 확정해야 합니다. "
    "특히 깊은 구간은 10 K 를 만들기 위한 펄스 에너지가 수십~수백 µJ 수준입니다(2.1절 표).",
    "선형 모델의 유효 범위는 표면 온도 상승 약 10 K 이내입니다. 50 K 민감도 값은 물성의 온도 의존성을 무시한 외삽입니다.",
    "3~4.6절의 검출 맵은 void 가 기둥 축 위에 있다고 가정한 결과입니다. 축에서 벗어난 void 는 같은 크기라도 신호가 훨씬 작으므로(4.7절), "
    "맵의 '검출 가능' 영역은 가장 유리한 위치 기준입니다.",
    "펌프는 반경 10 µm Gaussian 빔 기준입니다. 빔 크기나 형상(flat-top)을 바꾸면 결과가 달라집니다.",
    "void 형상은 회전타원체, 두께 10 µm 기준입니다. 경계 근처에서는 두께와 형상이 판정을 바꿀 수 있습니다(4.5절).",
    "비율 조건은 '같은 시각의 baseline 대비' 로 정의했습니다. 기준 시료(void 없는 TGV)와의 비교 측정이 가능하다는 전제입니다.",
])

doc.add_heading("부록 A. 주 격자 전체 값", level=1)
p("각 칸: 비율 [%] / 스케일된 ΔR/R 차이 [×10⁻⁶]. 파란 칸은 기본 기준에서 검출 가능.", size=9.5, color=(0x55, 0x55, 0x55))
table(["깊이 \\ 반폭 [µm]"] + ["%g" % w if w <= 40 else "80 (전체)" for w in WIDTHS],
      [["%g" % d] + ["%s / %s" % (pct(MAIN[(d, w)]["ratio"]), ("%.3g" if MAIN[(d, w)]["scaled_dRR_diff"] >= 1e-6 else "%.2g") % (MAIN[(d, w)]["scaled_dRR_diff"] * 1e6)) for w in WIDTHS] for d in DEPTHS],
      font_size=8.5, widths_cm=[2.4] + [2.0] * len(WIDTHS),
      fills=[[None] + ["DCE8F5" if MAIN[(d, w)]["detectable"] else "FDE9D2" for w in WIDTHS] for d in DEPTHS])

doc.add_heading("부록 B. 산출물", level=1)
bullets([
    "results/detection_map/cases.csv: 케이스별 한 행 (%d행). group 열: main(주 격자), boundary(경계용 추가 반폭), thickness, pulse." % len(ROWS),
    "results/detection_map/fig1~fig6 *.png: 2400×1350 px (200 dpi, 16:9).",
    "results/detection_map/summary.md: 깊이별 최소 반폭 요약표. summary.json: 같은 내용과 조건.",
    "results/detection_map/offaxis_cases.csv: 축에서 벗어난 void 스윕 (%d행). fig7_offaxis.png: 그 그림." % len(OFF),
    "studies/detection_map.py (스윕과 판정), detection_offaxis.py (축에서 벗어난 void, 3차원), detection_map_report.py (그림), "
    "detection_map_docx.py (이 보고서). 기존 시뮬레이터 코드는 수정하지 않았습니다.",
])
doc.add_heading("CSV 열 설명", level=2)
table(["열", "의미"],
      [["depth_um, half_width_um, thickness_um", "void 윗면 깊이, 가로 반폭, 세로 두께 [µm]"],
       ["blocks_rod", "반폭이 기둥 반경(40 µm)보다 큼 = 단면 전체 차단"],
       ["preset, tau_factor, tau_p_s, energy_J", "깊이 구간, 펄스폭 배수, 펄스폭 [s], 계산에 쓴 펄스 에너지 [J]"],
       ["window_s, window_factor, peak_at_window_end", "관측 시간창 [s], 자동 연장 배수, 연장 후에도 피크가 창 끝인지"],
       ["dT_base_max_K, dT_base_center_max_K", "baseline 피크: 프로브 가중 / 빔 중심 [K]"],
       ["void_minus_base_K, t_peak_s, base_at_peak_K", "void − baseline 피크 [K], 그 시각 [s], 그 시각의 baseline [K]"],
       ["dRR_diff, ratio", "ΔR/R 차이(스케일 전), 비율 (소수)"],
       ["scale, scaled_void_minus_base_K, scaled_dRR_diff", "스케일 계수 s (10 K 기준), 스케일된 온도 차이 [K], 스케일된 |ΔR/R 차이|"],
       ["noise_ok, ratio_ok, detectable, limiting", "기본 기준의 잡음 조건, 비율 조건, 최종 판정, 걸린 조건"]],
      widths_cm=[7.0, 9.4], align_right_from=9, font_size=9)

out = OUT / "ns-TTR_void_검출한계_시뮬레이션_결과보고.docx"
doc.save(out)
print("saved", out)
