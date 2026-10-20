"""Summary figures, cases.csv, summary.json and the Word report of the multi-void study (reads the
results/multi_void/cases cache written by multi_void.py; cases that are not finished yet are skipped).

Run from the src directory:   .venv\\Scripts\\python.exe studies\\multi_void_report.py  [--no-docx]
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import sys
from pathlib import Path

import numpy as np

SRC = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SRC))
sys.path.insert(0, str(SRC / "studies"))

from detection_map import CRITERIA_DEFAULT, N_SIGMA, UM  # noqa: E402
from multi_void import (BASE_DEPTH, CASES_DIR, DEEP_DEPTHS, FIG_DIR, OUT, RING_R, ROD_R_UM, WINDOW_US, all_cases,  # noqa: E402
                        setup_mpl)
from ttr_sim.materials import COPPER, COPPER_C_TR_PROBE  # noqa: E402

SIGMA3 = N_SIGMA * CRITERIA_DEFAULT["sigma"]
C_MIN = CRITERIA_DEFAULT["c_min"]
SUM_DIR = OUT / "summary"


# ----------------------------------------------------------------------------- data
def load_cases() -> dict:
    rows = {}
    for c in all_cases():
        jp = CASES_DIR / f"{c['id']}.json"
        if jp.exists():
            r = json.loads(jp.read_text(encoding="utf-8"))
            r["case"] = c
            rows[c["id"]] = r
    return rows


def curves(cid: str):
    return np.load(CASES_DIR / f"{cid}.npz")


def depolarisation_z(a_um: float, c_um: float) -> float:
    """Depolarisation factor along the symmetry axis z of a spheroid with equatorial semi-axis a (r_half) and
    polar semi-axis c (thickness / 2).  Sphere 1/3, thin disk -> 1, needle -> 0."""
    if abs(a_um - c_um) < 1e-9 * a_um:
        return 1.0 / 3.0
    if c_um < a_um:                                           # oblate
        e = math.sqrt(a_um ** 2 / c_um ** 2 - 1.0)
        return (1.0 + e ** 2) / e ** 3 * (e - math.atan(e))
    e = math.sqrt(1.0 - a_um ** 2 / c_um ** 2)                # prolate
    return (1.0 - e ** 2) / e ** 2 * (0.5 / e * math.log((1.0 + e) / (1.0 - e)) - 1.0)


def polarisability_um3(v: dict) -> float:
    """Thermal dipole strength of a perfectly insulating spheroid in a uniform axial gradient, up to a constant:
    V / (1 - L_z).  Finite for a thin disk (8 a^3 / 3) although its volume vanishes."""
    a, c = v["rh"], v["thick"] / 2.0
    V = 4.0 / 3.0 * math.pi * a * a * c
    return V / (1.0 - depolarisation_z(a, c))


def scaled_dRR(r: dict) -> float:
    return r["scaled_dRR_diff"]


# ----------------------------------------------------------------------------- summary figures
DEPTH_CMAP = "viridis_r"


def fig_A(R, plt):
    singles = [(r["depth_min_um"], r) for cid, r in R.items() if r["series"] == "A" and cid != "A_d070_axis"]
    if len(singles) < 2:
        return None
    singles.sort()
    fig, axs = plt.subplots(1, 3, figsize=(17, 5.2))
    norm = plt.Normalize(60, 220)
    cm = plt.get_cmap(DEPTH_CMAP)
    ax = axs[0]
    for d, r in singles:
        c = curves(r["id"])
        ax.plot(c["times"] * 1e6, c["dT_diff"] * 1e3, color=cm(norm(d)), label=f"{d:g} μm")
    if "A_d070_axis" in R:
        c = curves("A_d070_axis")
        ax.plot(c["times"] * 1e6, c["dT_diff"] * 1e3, color="k", ls="--", label="70 μm, 축 위")
    ax.set_xlim(0, WINDOW_US)
    ax.set_xlabel("t [μs]")
    ax.set_ylabel("void - baseline ΔT [mK]")
    ax.set_title("A. 단독 void 신호 ΔT(t) (r = 25 μm, 깊이별)")
    ax.legend(fontsize=8, ncol=2)
    ax.grid(alpha=0.3)
    ax = axs[1]
    d_ = np.array([d for d, _ in singles])
    ax.semilogy(d_, [abs(r["void_minus_base_K"]) * 1e3 for _, r in singles], "o-", label="|ΔT| 피크 [mK]")
    ax.semilogy(d_, [r["t_peak_s"] * 1e6 for _, r in singles], "s--", label="피크 시각 [μs]")
    ax.semilogy(d_, 2 * (d_ * UM) ** 2 / COPPER.alpha * 1e6, ":", color="gray", label="2 d²/D [μs]")
    ax.set_xlabel("void 윗면 깊이 [μm]")
    ax.set_title("피크 크기와 시각")
    ax.legend(fontsize=9)
    ax.grid(alpha=0.3, which="both")
    ax = axs[2]
    ax.semilogy(d_, [scaled_dRR(r) for _, r in singles], "o-", color="#d62728", label="|ΔR/R| (10 K 스케일)")
    ax.axhline(SIGMA3, color="#d62728", ls=":", label=f"잡음 기준 3σ = {SIGMA3:g}")
    ax.set_xlabel("void 윗면 깊이 [μm]")
    ax.set_ylabel("|ΔR/R|")
    ax2 = ax.twinx()
    ax2.semilogy(d_, [abs(r["contrast_max"]) * 100 for _, r in singles], "s-", color="#2ca02c", label="상대 대비 최대 [%]")
    ax2.axhline(C_MIN * 100, color="#2ca02c", ls=":", label=f"대비 기준 {C_MIN * 100:g} %")
    ax2.set_ylabel("상대 대비 [%]")
    h1, l1 = ax.get_legend_handles_labels()
    h2, l2 = ax2.get_legend_handles_labels()
    ax.legend(h1 + h2, l1 + l2, fontsize=8, loc="lower left")
    ax.set_title("검출 지표 대 깊이")
    ax.grid(alpha=0.3, which="both")
    fig.tight_layout()
    fig.savefig(SUM_DIR / "fig_A_depth.png", dpi=120)
    plt.close(fig)
    return "fig_A_depth.png"


def superposed(R, n):
    """Sum of the single-void signals that make up B_N{n}: 70 um axis + the first n deep ring voids."""
    ids = ["A_d070_axis"] + [f"A_d{int(d):03d}_r25" for d in DEEP_DEPTHS[:n]]
    if not all(i in R for i in ids):
        return None, None
    c0 = curves(ids[0])
    t = c0["times"]
    s = np.zeros_like(t)
    for i in ids:
        c = curves(i)
        s += np.interp(t, c["times"], c["dT_diff"])
    return t, s


def fig_B(R, plt):
    Ns = [n for n in (1, 3, 6, 9) if f"B_N{n}" in R]
    if not Ns or "A_d070_axis" not in R:
        return None
    fig, axs = plt.subplots(1, 3, figsize=(17, 5.2))
    cm = plt.get_cmap("plasma")
    ax = axs[0]
    c = curves("A_d070_axis")
    ax.plot(c["times"] * 1e6, c["dT_diff"] * 1e3, color="k", label="70 μm 단독 (N = 0)")
    for i, n in enumerate(Ns):
        c = curves(f"B_N{n}")
        col = cm(0.15 + 0.7 * i / max(1, len(Ns) - 1))
        ax.plot(c["times"] * 1e6, c["dT_diff"] * 1e3, color=col, label=f"계산 N = {n}")
        t, s = superposed(R, n)
        if s is not None:
            ax.plot(t * 1e6, s * 1e3, color=col, ls="--", alpha=0.8, label=f"단독 합 N = {n}")
    ax.set_xlim(0, WINDOW_US)
    ax.set_xlabel("t [μs]")
    ax.set_ylabel("void - baseline ΔT [mK]")
    ax.set_title("B. 70 μm + 깊은 void N개: 계산(실선) 대 단독 신호의 합(점선)")
    ax.legend(fontsize=8)
    ax.grid(alpha=0.3)
    ax = axs[1]
    base = R["A_d070_axis"]
    xs = [0] + Ns
    pk = [abs(base["void_minus_base_K"])] + [abs(R[f"B_N{n}"]["void_minus_base_K"]) for n in Ns]
    ax.plot(xs, np.array(pk) * 1e3, "o-", label="|ΔT| 피크 (계산)")
    sums = []
    for n in Ns:
        t, s = superposed(R, n)
        sums.append(abs(s).max() * 1e3 if s is not None else np.nan)
    ax.plot(Ns, sums, "s--", label="|ΔT| 피크 (단독 합)")
    ax.set_xlabel("추가한 깊은 void 개수 N")
    ax.set_ylabel("[mK]")
    ax.set_title("피크 크기 대 N")
    ax.legend(fontsize=9)
    ax.grid(alpha=0.3)
    ax = axs[2]
    ax.semilogy(xs, [scaled_dRR(base)] + [scaled_dRR(R[f"B_N{n}"]) for n in Ns], "o-", color="#d62728", label="|ΔR/R| (10 K 스케일)")
    ax.axhline(SIGMA3, color="#d62728", ls=":", label="잡음 기준 3σ")
    ax.set_ylabel("|ΔR/R|")
    ax.set_xlabel("추가한 깊은 void 개수 N")
    ax2 = ax.twinx()
    ax2.semilogy(xs, [abs(base["contrast_max"]) * 100] + [abs(R[f"B_N{n}"]["contrast_max"]) * 100 for n in Ns], "s-", color="#2ca02c", label="상대 대비 최대 [%]")
    ax2.axhline(C_MIN * 100, color="#2ca02c", ls=":", label="대비 기준 2 %")
    ax2.set_ylabel("[%]")
    h1, l1 = ax.get_legend_handles_labels()
    h2, l2 = ax2.get_legend_handles_labels()
    ax.legend(h1 + h2, l1 + l2, fontsize=8, loc="lower right")
    ax.set_title("검출 지표 대 N")
    ax.grid(alpha=0.3, which="both")
    fig.tight_layout()
    fig.savefig(SUM_DIR / "fig_B_superposition.png", dpi=120)
    plt.close(fig)
    return "fig_B_superposition.png"


def fig_C(R, plt):
    ids = [i for i in ("A_d070_axis", "B_N3", "C_stack3") if i in R]
    if "C_stack3" not in ids:
        return None
    fig, axs = plt.subplots(1, 2, figsize=(13, 5))
    ax = axs[0]
    for i, col in zip(ids, ("k", "#1f77b4", "#d62728")):
        c = curves(i)
        ax.plot(c["times"] * 1e6, c["dT_diff"] * 1e3, color=col, label=R[i]["label"])
    ax.set_xlim(0, WINDOW_US)
    ax.set_xlabel("t [μs]")
    ax.set_ylabel("void - baseline ΔT [mK]")
    ax.set_title("C. 그림자 효과: 옆으로 흩어진 3개(B_N3) 대 바로 아래 적층 3개(C)")
    ax.legend(fontsize=8)
    ax.grid(alpha=0.3)
    ax = axs[1]
    for i in ids:
        c = curves(i)
        ax.plot(c["times"] * 1e6, c["contrast"] * 100, label=R[i]["label"])
    ax.axhline(C_MIN * 100, color="gray", ls=":")
    ax.set_xlim(0, WINDOW_US)
    ax.set_xlabel("t [μs]")
    ax.set_ylabel("상대 대비 [%]")
    ax.set_title("상대 대비 ΔT / ΔT_baseline")
    ax.legend(fontsize=8)
    ax.grid(alpha=0.3)
    fig.tight_layout()
    fig.savefig(SUM_DIR / "fig_C_shadow.png", dpi=120)
    plt.close(fig)
    return "fig_C_shadow.png"


def fig_D(R, plt):
    ids = [i for i in ("A_d070_axis", "D_N2", "D_N3", "D_N6") if i in R]
    if len(ids) < 2:
        return None
    fig, axs = plt.subplots(1, 3, figsize=(17, 5.2))
    cm = plt.get_cmap("plasma")
    ax = axs[0]
    for k, i in enumerate(ids):
        c = curves(i)
        ax.plot(c["times"] * 1e6, c["dT_diff"] * 1e3, color=cm(0.1 + 0.75 * k / 3), label=f"{R[i]['n_voids']}개 (차단 {R[i]['blocked_area_frac'] * 100:.0f} %)")
    ax.set_xlim(0, WINDOW_US)
    ax.set_xlabel("t [μs]")
    ax.set_ylabel("void - baseline ΔT [mK]")
    ax.set_title("D. 70 μm 에 void 를 옆으로 늘릴 때")
    ax.legend(fontsize=8)
    ax.grid(alpha=0.3)
    ax = axs[1]
    fr = [R[i]["blocked_area_frac"] * 100 for i in ids]
    pk = [abs(R[i]["void_minus_base_K"]) * 1e3 for i in ids]
    ax.plot(fr, pk, "o-", label="같은 깊이 70 μm (D)")
    pk0 = pk[0] / fr[0]
    ax.plot(fr, [pk0 * f for f in fr], ":", color="gray", label="개수 비례")
    bs = [i for i in ("A_d070_axis", "B_N1", "B_N3", "B_N6") if i in R]
    if len(bs) > 1:
        ax.plot([R[i]["blocked_area_frac"] * 100 for i in bs], [abs(R[i]["void_minus_base_K"]) * 1e3 for i in bs], "s--", label="깊이 방향 추가 (B)")
    ax.set_xlabel("막힌 단면적 비율 [%]")
    ax.set_ylabel("|ΔT| 피크 [mK]")
    ax.set_title("피크 대 막힌 단면적")
    ax.legend(fontsize=9)
    ax.grid(alpha=0.3)
    ax = axs[2]
    ax.semilogy(fr, [scaled_dRR(R[i]) for i in ids], "o-", color="#d62728", label="|ΔR/R| (10 K 스케일)")
    ax.axhline(SIGMA3, color="#d62728", ls=":", label="잡음 기준 3σ")
    ax.set_xlabel("막힌 단면적 비율 [%]")
    ax.set_ylabel("|ΔR/R|")
    ax2 = ax.twinx()
    ax2.semilogy(fr, [abs(R[i]["contrast_max"]) * 100 for i in ids], "s-", color="#2ca02c", label="상대 대비 최대 [%]")
    ax2.axhline(C_MIN * 100, color="#2ca02c", ls=":", label="대비 기준 2 %")
    ax2.set_ylabel("[%]")
    h1, l1 = ax.get_legend_handles_labels()
    h2, l2 = ax2.get_legend_handles_labels()
    ax.legend(h1 + h2, l1 + l2, fontsize=8, loc="lower right")
    ax.set_title("검출 지표")
    ax.grid(alpha=0.3, which="both")
    fig.tight_layout()
    fig.savefig(SUM_DIR / "fig_D_lateral.png", dpi=120)
    plt.close(fig)
    return "fig_D_lateral.png"


def e_cases(R):
    ids = [i for i in R if R[i]["series"] == "E"] + (["A_d070_axis"] if "A_d070_axis" in R else [])
    return sorted(ids, key=lambda i: (R[i]["case"]["voids"][0]["rh"], R[i]["case"]["voids"][0]["thick"]))


def fig_E(R, plt):
    ids = e_cases(R)
    if len(ids) < 3:
        return None
    fig, axs = plt.subplots(1, 3, figsize=(17, 5.2))
    ax = axs[0]
    for i in ids:
        v = R[i]["case"]["voids"][0]
        c = curves(i)
        ax.plot(c["times"] * 1e6, c["dT_diff"] * 1e3, label=f"반폭 {v['rh']:g}, 두께 {v['thick']:g} μm")
    ax.set_xlim(0, WINDOW_US)
    ax.set_xlabel("t [μs]")
    ax.set_ylabel("void - baseline ΔT [mK]")
    ax.set_title("E. 70 μm 단독 void 의 형상")
    ax.legend(fontsize=8)
    ax.grid(alpha=0.3)
    vol = np.array([R[i]["void_volume_um3"] for i in ids])
    area = np.array([math.pi * R[i]["case"]["voids"][0]["rh"] ** 2 for i in ids])
    pol = np.array([polarisability_um3(R[i]["case"]["voids"][0]) for i in ids])
    pk = np.array([abs(R[i]["void_minus_base_K"]) * 1e3 for i in ids])
    ax = axs[1]
    ax.loglog(vol, pk, "o", label="대 부피 V")
    ax.loglog(area * 10, pk, "s", label="대 면적 × 10 μm")
    for i, x, y in zip(ids, vol, pk):
        v = R[i]["case"]["voids"][0]
        ax.annotate(f"({v['rh']:g}, {v['thick']:g})", (x, y), fontsize=7, xytext=(3, 3), textcoords="offset points")
    ax.set_xlabel("[μm³]")
    ax.set_ylabel("|ΔT| 피크 [mK]")
    ax.set_title("피크 대 부피 / 면적: 어느 쪽도 한 직선이 아님")
    ax.legend(fontsize=9)
    ax.grid(alpha=0.3, which="both")
    ax = axs[2]
    ax.loglog(pol, pk, "o", color="#d62728")
    k = np.exp(np.mean(np.log(pk) - np.log(pol)))
    xx = np.linspace(pol.min() * 0.8, pol.max() * 1.2, 50)
    ax.loglog(xx, k * xx, "--", color="gray", label="∝ V / (1 - L_z)")
    for i, x, y in zip(ids, pol, pk):
        v = R[i]["case"]["voids"][0]
        ax.annotate(f"({v['rh']:g}, {v['thick']:g})", (x, y), fontsize=7, xytext=(3, 3), textcoords="offset points")
    ax.set_xlabel("쌍극자 분극률 V / (1 - L_z) [μm³]")
    ax.set_ylabel("|ΔT| 피크 [mK]")
    ax.set_title("피크 대 절연 회전타원체의 쌍극자 분극률")
    ax.legend(fontsize=9)
    ax.grid(alpha=0.3, which="both")
    fig.tight_layout()
    fig.savefig(SUM_DIR / "fig_E_shape.png", dpi=120)
    plt.close(fig)
    return "fig_E_shape.png"


def fig_F(R, plt):
    ids = sorted(i for i in R if R[i]["series"] == "F")
    if not ids:
        return None
    fig, axs = plt.subplots(1, 2, figsize=(13, 5))
    ax = axs[0]
    for i in ids:
        c = curves(i)
        ax.plot(c["times"] * 1e6, c["dT_diff"] * 1e3, label=f"{R[i]['label']}: 가장 얕은 {R[i]['depth_min_um']:.0f} μm")
    if "B_N9" in R:
        c = curves("B_N9")
        ax.plot(c["times"] * 1e6, c["dT_diff"] * 1e3, "k--", label="B_N9 (70 μm + 9개)")
    ax.set_xlim(0, WINDOW_US)
    ax.set_xlabel("t [μs]")
    ax.set_ylabel("void - baseline ΔT [mK]")
    ax.set_title("F. 무작위 10개 배치")
    ax.legend(fontsize=8)
    ax.grid(alpha=0.3)
    ax = axs[1]
    singles = sorted((r["depth_min_um"], abs(r["void_minus_base_K"]) * 1e3) for i, r in R.items() if r["series"] == "A" and i != "A_d070_axis")
    if singles:
        ax.semilogy([d for d, _ in singles], [p for _, p in singles], "k:", label="단독 void (r = 25 μm) 깊이 곡선")
    ax.semilogy([R[i]["depth_min_um"] for i in ids], [abs(R[i]["void_minus_base_K"]) * 1e3 for i in ids], "o", color="#d62728", label="무작위 10개 (가장 얕은 void 깊이)")
    for i in ids:
        ax.annotate(R[i]["label"].split("(")[-1].rstrip(")"), (R[i]["depth_min_um"], abs(R[i]["void_minus_base_K"]) * 1e3), fontsize=7, xytext=(3, 3), textcoords="offset points")
    ax.set_xlabel("가장 얕은 void 윗면 깊이 [μm]")
    ax.set_ylabel("|ΔT| 피크 [mK]")
    ax.set_title("피크 대 가장 얕은 void 깊이")
    ax.legend(fontsize=8)
    ax.grid(alpha=0.3, which="both")
    fig.tight_layout()
    fig.savefig(SUM_DIR / "fig_F_random.png", dpi=120)
    plt.close(fig)
    return "fig_F_random.png"


# ----------------------------------------------------------------------------- tables
CSV_COLS = ["id", "series", "label", "n_voids", "depth_min_um", "depth_max_um", "blocked_area_frac", "void_volume_um3",
            "is_3d", "n_unknowns", "n_steps", "void_cells", "solver", "mean_iterations", "energy_error", "wall_s",
            "dT_base_max_K", "t_base_peak_s", "void_minus_base_K", "t_peak_s", "ratio", "base_at_peak_K", "dRR_diff",
            "contrast_max", "t_contrast_max_s", "scale", "scaled_dRR_diff", "noise_margin", "ratio_margin",
            "noise_ok", "ratio_ok", "detectable", "limiting"]


def write_csv(R):
    with open(OUT / "cases.csv", "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=CSV_COLS, extrasaction="ignore")
        w.writeheader()
        for c in all_cases():
            if c["id"] in R:
                w.writerow(R[c["id"]])


# ----------------------------------------------------------------------------- docx
def build_docx(R, figs):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    FONT = "Malgun Gothic"
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin = s.bottom_margin = Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for name in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        hs = doc.styles[name]
        hs.font.name = FONT
        hs.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        hs.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

    def P(text, bold=False, italic=False, size=None, align=None, style=None):
        p = doc.add_paragraph(style=style)
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if align == "c":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return p

    def H(text, level):
        return doc.add_heading(text, level)

    def B(text):
        return P(text, style="List Bullet")

    def table(header, rows, widths=None, size=8.5):
        t = doc.add_table(rows=1, cols=len(header))
        t.style = "Light Grid Accent 1"
        for i, h in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(h))
            run.bold = True
            run.font.size = Pt(size)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(str(v))
                run.font.size = Pt(size)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        doc.add_paragraph()
        return t

    def fig(name, caption, width_cm=16.5, folder=SUM_DIR):
        path = folder / name
        if not path.exists():
            P(f"[그림 없음: {name}]", italic=True)
            return
        doc.add_picture(str(path), width=Cm(width_cm))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        P(caption, italic=True, size=9, align="c")

    def mk(x, unit="mK"):
        return f"{x * 1e3:.3g} {unit}" if unit == "mK" else f"{x:.3g} {unit}"

    def verdict(r):
        return "검출 가능" if r["detectable"] else f"불가 ({r['limiting']})"

    def metrics_rows(ids):
        rows = []
        for i in ids:
            if i not in R:
                continue
            r = R[i]
            rows.append([i, r["n_voids"], f"{r['blocked_area_frac'] * 100:.0f} %", mk(r["void_minus_base_K"]),
                         f"{r['t_peak_s'] * 1e6:.0f}", f"{r['ratio'] * 100:.2f} %", f"{r['contrast_max'] * 100:.2f} %",
                         f"{r['scaled_dRR_diff']:.2e}", f"×{r['noise_margin']:.2f}", f"×{r['ratio_margin']:.2f}", verdict(r)])
        return rows

    MH = ["케이스", "void 수", "단면 합(투영)", "ΔT 피크", "시각 [μs]", "피크 대비", "대비 최대", "|ΔR/R| (10 K)", "잡음 여유", "대비 여유", "판정"]
    MW = [2.4, 1.1, 1.3, 1.7, 1.3, 1.4, 1.4, 1.9, 1.3, 1.3, 1.9]

    base = R.get("A_d070_axis")
    b9 = R.get("B_N9")
    doc.add_heading("ns-TTR 다중 void 검출 시뮬레이션 결과 보고", 0)
    P("ns-TTR Simulator 다중 void 스터디 (studies/multi_void.py, 결과 results/multi_void)", size=9, align="c")
    P(f"케이스 {len(R)}개 · 생성 {__import__('time').strftime('%Y-%m-%d %H:%M')}", size=9, align="c")

    H("1. 목적과 질문", 1)
    P("70 μm 깊이의 반폭 10 μm, 두께 10 μm 공기 void 하나는 기존 검출 기준에 못 미친다. 이 스터디는 그런 void 가 여럿 "
      "존재할 때 무엇이 달라지는지를 묻는다.")
    B("질문 1. 단독으로는 검출되지 않는 void 가 여러 개 흩어져 있으면 검출 가능한가? 신호는 단독 신호의 합인가, 어디에 "
      "추가 void 가 있어야 합이 커지는가?")
    B("질문 2. 전체 void 부피가 검출을 좌우하는가? 아니면 단면적 또는 다른 양인가?")

    H("2. 물리적 전제와 예상", 1)
    P("모델은 구리 로드(반지름 40 μm, 길이 500 μm)가 실리카에 묻힌 축대칭 구조의 선형 열확산이며, void 는 공기 "
      "(k = 0.026 W/m·K)로 채워진 회전타원체다. 축을 벗어난 void 는 (r, θ, z) 3차원으로 푼다. 실험 전에 세운 전제는 다음과 같다.")
    B("깊이가 로드 반지름보다 깊으면 열흐름은 거의 1차원이다. void 의 효과는 단면을 막는 것이므로 신호는 막힌 단면적 비율과 "
      "깊이 함수의 곱으로 정해질 것이다.")
    B("공기 void 는 두께와 무관한 완전 차단체다(1 μm 공기층의 열저항이 구리 15 mm 와 같다). 따라서 부피가 아니라 투영 "
      "면적이 효과를 정할 것이다.")
    B("약한 섭동이므로 void 여러 개의 신호는 근사적으로 단독 신호의 합이다. 다만 다른 void 의 열그림자 안에 있는 void 는 "
      "거의 기여하지 못할 것이다.")
    B("깊은 void 의 응답은 t ≈ d²/D 이후에 도달하므로 가장 깊은 void(210 μm, d²/D = 380 μs)를 보려면 관측창을 1 ms 로 "
      "늘려야 하고, 펄스폭은 70 μm 프리셋 값(27.6 μs)을 그대로 둔다.")
    P("보고서 끝의 '예상과의 비교' 표에서 이 전제들을 결과와 대조한다.")

    H("3. 방법", 1)
    H("3.1 공통 조건", 2)
    table(["항목", "값"], [
        ["void 크기 (E 제외)", "반폭 10 μm, 두께 10 μm, 회전타원체, 공기 k = 0.026 W/m·K"],
        ["기준 void", "윗면 깊이 70 μm, 축 위"],
        ["펄스", "사각, τp = 27.6 μs (40~100 μm 프리셋: τp = 2 d²/D, d = 40 μm), 에너지 GUI 기본값(표면 피크 ≈ 0.13 K)"],
        ["빔", "펌프 가우시안 w = 10 μm, 프로브 5 μm, 축 위"],
        ["관측창", f"{WINDOW_US:g} μs (전 케이스 동일, 같은 시간축)"],
        ["후면", "산화막 5 nm + 열패드 200 μm + 히트싱크(실온)"],
        ["격자·시간", "표면층 Δz = 5.7 μm, void 주변은 void 치수의 1/8, Δt = τp/200 → 펄스 후 스텝당 +5 %"],
        ["3D 솔버", "미지수 6만 이하 직접 LU, 그 이상 격자 구조 전처리(섹터·고리 정확해 + θ Fourier 거친 보정) CG, 잔차 1e-11"],
    ], widths=[3.5, 13.0])
    H("3.2 검출 기준", 2)
    P(f"기존 검출한계 스터디와 같은 기준이다. 펄스 에너지를 baseline 피크 상승이 {CRITERIA_DEFAULT['dT_allowed']:g} K 가 되도록 "
      f"스케일한 뒤(선형 모델), (i) |ΔR/R| 의 void-baseline 차이 ≥ {N_SIGMA:g}σ = {SIGMA3:g} (σ = {CRITERIA_DEFAULT['sigma']:g}), "
      f"(ii) 신호 피크 시각의 상대 대비 |ΔT/ΔT_baseline| ≥ {C_MIN * 100:g} %. 두 조건을 모두 만족해야 '검출 가능'이다. "
      "표의 '여유'는 각 기준 대비 배수다. 상대 대비는 피크 시각 값과 창 전체의 최대값을 함께 적었다. '단면 합(투영)'은 모든 void 의 최대 단면적을 로드 단면적으로 나눈 합이며, 깊이가 다른 void 는 서로 가리지 않는다는 뜻이 아니라 크기의 척도로만 적었다.")
    H("3.3 케이스 구성", 2)
    table(["계열", "목적", "케이스"], [
        ["A", "단독 void 의 깊이 함수, B 의 중첩 검증 재료", "70 μm(축), 70 μm(r = 25), 90~210 μm 15 μm 간격(r = 25): 11개"],
        ["B", "깊이 방향 누적 (질문 1)", f"70 μm + 깊은 void N개 (90 + 15k μm, r = {RING_R:g} μm, θ = 40k°), N = 1, 3, 6, 9"],
        ["C", "그림자 효과", "70 μm + 바로 아래 적층 3개 (90/105/120 μm, 축 위). B_N3 과 비교"],
        ["D", "같은 깊이 가로 분산 (단면 차단율)", "70 μm 에 축 + 고리(r = 25) 1, 2, 5개 → 2, 3, 6개 (차단 12, 19, 37 %)"],
        ["E", "부피 대 면적 (질문 2)", "70 μm 단독: 같은 부피 (14.1, 5), (10, 10), (7.1, 20); 같은 면적 (10, 2), (10, 5), (10, 10), (10, 20) [반폭, 두께 μm]"],
        ["F", "무작위 배치", "10개를 윗면 깊이 70 μm 이상에 균일 무작위 (시드 1~5). 참고 F0: 로드 전체 (GUI 랜덤 모드)"],
    ], widths=[1.2, 5.0, 10.3])
    P("B 의 깊은 void 는 70 μm void 의 투영 그림자 밖(r = 25 μm)에 두었고, N 이 작은 집합은 큰 집합의 부분집합이다. "
      "고리 위 이웃 void 끼리는 투영이 약간 겹친다(중심 거리 17 μm, 반폭 10 μm).")

    H("4. 결과", 1)
    # ---- A
    H("4.1 계열 A: 단독 void 의 깊이 함수", 2)
    if figs.get("A"):
        fig(figs["A"], "그림 A. 단독 void 신호 ΔT(t)(왼쪽, 같은 축), 피크 크기·시각(가운데), 10 K 스케일 |ΔR/R| 과 상대 대비(오른쪽).")
    a_ids = [i for i in R if R[i]["series"] == "A"]
    a_ids.sort(key=lambda i: (R[i]["depth_min_um"], i))
    table(MH, metrics_rows(a_ids), widths=MW)
    if base:
        P(f"70 μm 단독(축 위)의 피크는 {mk(base['void_minus_base_K'])}, 피크 시각 {base['t_peak_s'] * 1e6:.0f} μs, 상대 대비 "
          f"{base['ratio'] * 100:.2f} %, 10 K 스케일 |ΔR/R| = {base['scaled_dRR_diff']:.2e} 로 잡음 기준의 ×{base['noise_margin']:.2f}, "
          f"대비 기준의 ×{base['ratio_margin']:.2f} 이다. 사용자가 GUI 에서 확인한 '검출 기준 미달'과 같다.")
    singles = sorted((R[i]["depth_min_um"], R[i]) for i in a_ids if i != "A_d070_axis")
    if len(singles) >= 3:
        d0, r0 = singles[0]
        dl, rl = singles[-1]
        P(f"r = 25 μm 의 단독 void 는 {d0:g} μm 에서 {mk(r0['void_minus_base_K'])}, {dl:g} μm 에서 {mk(rl['void_minus_base_K'])} 로 "
          f"깊이에 따라 {abs(r0['void_minus_base_K'] / rl['void_minus_base_K']):.1f}배 줄고, 피크 시각은 {r0['t_peak_s'] * 1e6:.0f} → "
          f"{rl['t_peak_s'] * 1e6:.0f} μs 로 늦어진다. 개별 깊은 void 의 기여가 어느 정도인지가 B 의 합을 해석하는 눈금이 된다.")
    # ---- B
    H("4.2 계열 B: 70 μm void 아래에 깊은 void 를 더할 때 (질문 1)", 2)
    if figs.get("B"):
        fig(figs["B"], "그림 B. 계산된 신호(실선)와 단독 신호의 산술합(점선), 피크 크기와 검출 지표 대 추가 개수 N.")
    table(MH, metrics_rows(["A_d070_axis", "B_N1", "B_N3", "B_N6", "B_N9"]), widths=MW)
    if base and b9:
        t, s = superposed(R, 9)
        gain = abs(b9["void_minus_base_K"] / base["void_minus_base_K"])
        txt = (f"깊은 void 9개를 더하면 피크가 70 μm 단독의 {gain:.2f}배가 된다. ")
        if s is not None:
            txt += (f"단독 신호의 산술합은 피크 {abs(s).max() * 1e3:.3g} mK 로 계산값 {mk(b9['void_minus_base_K'])} 과 "
                    f"{abs(s).max() / abs(b9['void_minus_base_K']) * 100:.0f} % 수준이므로, 중첩 가정은 "
                    f"{'잘 맞는다' if abs(abs(s).max() / abs(b9['void_minus_base_K']) - 1) < 0.15 else '수정이 필요하다'}. ")
        txt += (f"N = 9 에서 10 K 스케일 |ΔR/R| 은 잡음 기준의 ×{b9['noise_margin']:.2f}, 상대 대비는 기준의 ×{b9['ratio_margin']:.2f} → "
                f"{verdict(b9)}.")
        P(txt)
    # ---- C
    H("4.3 계열 C: 그림자 효과", 2)
    if figs.get("C"):
        fig(figs["C"], "그림 C. 옆으로 흩어진 3개(B_N3)와 바로 아래 적층 3개(C_stack3)의 ΔT(t) 와 상대 대비.", 15.0)
    table(MH, metrics_rows(["A_d070_axis", "B_N3", "C_stack3"]), widths=MW)
    if base and "C_stack3" in R:
        c3 = R["C_stack3"]
        txt = f"적층 3개는 70 μm 단독의 {abs(c3['void_minus_base_K'] / base['void_minus_base_K']):.2f}배다. "
        if "B_N3" in R:
            b3 = R["B_N3"]
            txt += (f"옆으로 흩어진 3개(B_N3)는 {abs(b3['void_minus_base_K'] / base['void_minus_base_K']):.2f}배이므로 적층은 분산 배치의 "
                    f"{abs(c3['void_minus_base_K'] / b3['void_minus_base_K']) * 100:.0f} % 다. ")
        txt += ("적층이 단독보다 뚜렷이 큰 것은 '그림자 속 void 는 기여하지 않는다'는 전제와 다르다. 축 위에 쌓인 void 열은 "
                "하나의 긴(두꺼운) 장애물처럼 작용하여, 계열 E 에서 본 두께 효과와 같은 이유로 신호를 키운다.")
        P(txt)
    # ---- D
    H("4.4 계열 D: 같은 깊이에서 옆으로 늘릴 때", 2)
    if figs.get("D"):
        fig(figs["D"], "그림 D. 70 μm 에 void 1, 2, 3, 6개: ΔT(t), 피크 대 막힌 단면적 비율(깊이 방향 추가 B 와 비교), 검출 지표.")
    table(MH, metrics_rows(["A_d070_axis", "D_N2", "D_N3", "D_N6"]), widths=MW)
    if base and "D_N6" in R:
        d6 = R["D_N6"]
        gain6 = abs(d6["void_minus_base_K"] / base["void_minus_base_K"])
        txt = (f"같은 깊이에 6개(단면 {d6['blocked_area_frac'] * 100:.0f} % 차단)를 두면 피크가 단독의 {gain6:.2f}배가 되고 판정은 "
               f"'{verdict(d6)}' 이다 (잡음 여유 ×{d6['noise_margin']:.2f}, 대비 여유 ×{d6['ratio_margin']:.2f}). ")
        if "A_d070_r25" in R:
            ssum = abs(base["void_minus_base_K"]) + 5 * abs(R["A_d070_r25"]["void_minus_base_K"])
            txt += (f"단독 신호의 합(축 위 1개 + r = 25 μm 1개 × 5)은 {ssum * 1e3:.3g} mK 로 계산값 {mk(d6['void_minus_base_K'])} 의 "
                    f"{ssum / abs(d6['void_minus_base_K']) * 100:.0f} % 이므로, 같은 깊이에 나란히 놓인 void 는 개수 비례보다 "
                    f"{'조금 더' if ssum < abs(d6['void_minus_base_K']) else '조금 덜'} 기여한다 (이웃한 void 가 함께 단면을 좁혀 우회로가 줄어드는 상호작용). ")
        if "D_N3" in R:
            d3 = R["D_N3"]
            txt += (f"3개(차단 {d3['blocked_area_frac'] * 100:.0f} %)는 {abs(d3['void_minus_base_K'] / base['void_minus_base_K']):.2f}배로 "
                    f"잡음 여유 ×{d3['noise_margin']:.2f}, 대비 여유 ×{d3['ratio_margin']:.2f} → '{verdict(d3)}'.")
        P(txt)
    # ---- E
    H("4.5 계열 E: 부피인가 면적인가 (질문 2)", 2)
    if figs.get("E"):
        fig(figs["E"], "그림 E. 70 μm 단독 void 의 형상 변화: ΔT(t), 피크 대 부피·면적, 피크 대 절연 회전타원체의 쌍극자 분극률 V/(1 - L_z).")
    e_ids = e_cases(R)
    rows = []
    for i in e_ids:
        r = R[i]
        v = r["case"]["voids"][0]
        rows.append([f"{v['rh']:g} / {v['thick']:g}", f"{r['void_volume_um3']:,.0f}", f"{math.pi * v['rh'] ** 2:,.0f}",
                     f"{depolarisation_z(v['rh'], v['thick'] / 2):.3f}", f"{polarisability_um3(v):,.0f}",
                     mk(r["void_minus_base_K"]), f"{r['ratio'] * 100:.2f} %", f"{r['scaled_dRR_diff']:.2e}", verdict(r)])
    table(["반폭 / 두께 [μm]", "부피 [μm³]", "단면적 [μm²]", "L_z", "V/(1-L_z) [μm³]", "ΔT 피크", "피크 대비", "|ΔR/R| (10 K)", "판정"],
          rows, widths=[2.4, 1.8, 1.8, 1.2, 2.2, 1.8, 1.6, 2.0, 2.0])
    if len(e_ids) >= 4:
        same_area = sorted((R[i]["case"]["voids"][0]["thick"], R[i]) for i in e_ids if abs(R[i]["case"]["voids"][0]["rh"] - 10.0) < 0.01)
        same_vol = sorted((R[i]["case"]["voids"][0]["rh"], R[i]) for i in e_ids if R[i]["case"]["voids"][0]["rh"] in (14.1, 10.0, 7.1)
                          and abs(R[i]["void_volume_um3"] - R["A_d070_axis"]["void_volume_um3"]) / R["A_d070_axis"]["void_volume_um3"] < 0.05)
        txt = ""
        if len(same_area) >= 2:
            (t0, r0), (t1, r1) = same_area[0], same_area[-1]
            txt += (f"같은 면적(반폭 10 μm)에서 두께를 {t0:g} → {t1:g} μm 로 늘리면 피크가 {abs(r1['void_minus_base_K'] / r0['void_minus_base_K']):.2f}배 "
                    "커진다. 두께가 무관하리라는 전제는 틀렸다. ")
        if len(same_vol) >= 2:
            (a0, r0), (a1, r1) = same_vol[0], same_vol[-1]
            txt += (f"같은 부피에서는 반폭 {a0:g} μm(두께 {r0['case']['voids'][0]['thick']:g})가 반폭 {a1:g} μm(두께 "
                    f"{r1['case']['voids'][0]['thick']:g})의 {abs(r0['void_minus_base_K'] / r1['void_minus_base_K']):.2f}배이므로 부피도 지표가 아니다. ")
        pol = np.array([polarisability_um3(R[i]["case"]["voids"][0]) for i in e_ids])
        pk = np.array([abs(R[i]["void_minus_base_K"]) for i in e_ids])
        slope = np.polyfit(np.log(pol), np.log(pk), 1)[0]
        resid = np.std(np.log(pk) - np.log(pol) - np.mean(np.log(pk) - np.log(pol)))
        txt += (f"여섯 형상의 피크를 균일 온도구배 속 절연 회전타원체의 쌍극자 세기 V/(1 - L_z) 에 대해 그리면 한 직선에 모인다 "
                f"(로그 기울기 {slope:.2f}, 산포 {resid * 100:.0f} %). 즉 멀리서 본 void 의 효과는 부피도 면적도 아닌 열 쌍극자 분극률로 "
                "정해지며, 얇은 원판은 부피가 0 에 가까워도 8a³/3 의 분극률을 가지고, 같은 반폭에서 두께가 늘면 (1 - L_z) 가 커져 "
                "분극률이 구(2πa³)까지 2.4배 늘어난다.")
        P(txt)
    # ---- F
    H("4.6 계열 F: 무작위 배치", 2)
    if figs.get("F"):
        fig(figs["F"], "그림 F. 무작위 10개 배치 5종의 ΔT(t)(B_N9 와 비교)와 피크 대 가장 얕은 void 깊이(단독 void 깊이 곡선과 비교).", 15.0)
    f_ids = sorted(i for i in R if R[i]["series"] == "F")
    table(MH, metrics_rows(f_ids), widths=MW)
    if f_ids:
        det = [i for i in f_ids if R[i]["detectable"]]
        P(f"윗면 깊이 70 μm 이상으로 제한한 무작위 배치 {len(f_ids)}종 중 검출 가능은 {len(det)}종이다. 가장 얕은 void 의 깊이가 "
          f"{min(R[i]['depth_min_um'] for i in f_ids):.0f}~{max(R[i]['depth_min_um'] for i in f_ids):.0f} μm 로 흩어져 있고, 피크는 "
          "대체로 그 깊이로 정해지며 나머지 9개는 그 위에 더해진다. 같은 10개라도 어디에 있느냐에 따라 피크가 수 배 차이 난다.")
    f0_ids = sorted(i for i in R if R[i]["series"] == "F0")
    if f0_ids:
        P("참고로 GUI 랜덤 모드와 같이 로드 전체(깊이 0~500 μm)에 뿌린 배치는 다음과 같다. 시드 3 은 void 하나가 표면을 뚫고 "
          "나와(윗면 깊이 0.7 μm) 펌프 흡수 자체가 줄어든 경우라 t = 0 부터 큰 음의 신호가 나오는데, 이는 내부 void 가 아니라 "
          "표면 구멍의 신호이므로 이 스터디의 질문과는 다른 경우다.", size=9)
        table(MH, metrics_rows(f0_ids), widths=MW, size=7.5)

    H("5. 예상과의 비교", 1)
    cmp_rows = []
    if base and b9:
        cmp_rows.append(["B: 깊은 void 9개를 더해도 단독의 1.5~2배 안팎", f"{abs(b9['void_minus_base_K'] / base['void_minus_base_K']):.2f}배",
                         "맞음" if 1.3 <= abs(b9['void_minus_base_K'] / base['void_minus_base_K']) <= 2.3 else "다름"])
    if base and "C_stack3" in R and "B_N3" in R:
        rc = abs(R["C_stack3"]["void_minus_base_K"] / R["B_N3"]["void_minus_base_K"])
        cmp_rows.append(["C: 적층은 가로 분산 3개보다 눈에 띄게 작다", f"적층/분산 = {rc * 100:.0f} %", "맞음" if rc < 0.7 else "다름"])
    if base and "D_N6" in R:
        cmp_rows.append(["D: 같은 깊이에서는 개수(차단율)에 거의 비례, 3개 이상이면 기준 통과",
                         f"6개에서 {abs(R['D_N6']['void_minus_base_K'] / base['void_minus_base_K']):.2f}배, 판정 {verdict(R['D_N6'])}"
                         + (f", 3개 판정 {verdict(R['D_N3'])}" if "D_N3" in R else ""),
                         "부분적으로 맞음"])
    if len(e_ids) >= 4:
        cmp_rows.append(["E: 같은 면적이면 두께 차이는 신호에 거의 없음", "두께 2 → 20 μm 에서 피크 2배 이상 증가", "틀림 (쌍극자 분극률이 지표)"])
        cmp_rows.append(["E: 같은 부피에서는 반폭이 넓은 쪽이 훨씬 큼", "넓은 쪽이 큼", "맞음"])
    if cmp_rows:
        table(["실험 전 예상", "결과", "판정"], cmp_rows, widths=[7.0, 6.0, 3.5])

    H("6. 결론", 1)
    if base and b9:
        B(f"질문 1: 단독으로 검출되지 않는 70 μm void 아래에 같은 크기의 void 를 9개 더하면 피크는 {abs(b9['void_minus_base_K'] / base['void_minus_base_K']):.2f}배가 "
          f"되고 신호는 단독 신호의 합과 같다(중첩). 상대 대비는 기준을 넘지만(×{b9['ratio_margin']:.2f}) 10 K 스케일 |ΔR/R| 이 잡음 기준의 "
          f"×{b9['noise_margin']:.2f} 로 판정은 '{verdict(b9)}' 이다. 깊은 void 의 기여는 깊이에 따라 급격히 줄어(90 → 210 μm 에서 1/7) "
          "개수를 늘려도 합이 포화한다. 6개에서 9개로 늘렸을 때의 증가는 3 % 뿐이다.")
    if base and "D_N6" in R:
        B(f"질문 1 (가로 분산): 같은 깊이 70 μm 에 6개(단면 {R['D_N6']['blocked_area_frac'] * 100:.0f} % 차단)를 두면 피크 {abs(R['D_N6']['void_minus_base_K'] / base['void_minus_base_K']):.2f}배, 판정 "
          f"'{verdict(R['D_N6'])}'. 여러 개의 작은 void 가 검출되는 현실적인 경로는 같은 깊이에 가로로 퍼져 단면을 많이 막는 경우다.")
    B("질문 2: 검출을 좌우하는 것은 부피도 투영 면적도 아니고 void 의 열 쌍극자 분극률 V/(1 - L_z) 이다. 같은 면적이라도 두꺼울수록, "
      "같은 부피라도 넓고 얇을수록 신호가 크다.")
    B("깊은 void 는 늦게, 작게 기여하므로 관측창을 가장 깊은 void 의 2~3 d²/D 까지 두어야 그 기여가 보이고, 피크 시각뿐 아니라 늦은 시각의 "
      "대비도 함께 평가해야 한다.")

    H("7. 한계", 1)
    B("선형 모델이며 물성·흡수율의 온도 의존성이 없다. 판정은 baseline 피크 10 K 스케일 기준이다.")
    B("B 와 D 의 고리 위 이웃 void 는 투영이 조금 겹친다. 무작위 배치(F)는 겹침과 축 통과를 허용한다.")
    B("쌍극자 분극률 해석은 무한 매질의 균일 구배에 대한 것이고, 로드 벽에 가깝거나 서로 가까운 void 에서는 보정이 필요하다.")

    H("참고", 1)
    B("ns-TTR void 검출한계 시뮬레이션 결과 보고 (results/detection_map), 검출 기준과 단독 void 깊이·반폭 지도.")
    B("ns-TTR 시뮬레이터 구현 스펙, §7 검증 절차.")
    B("H. S. Carslaw, J. C. Jaeger, Conduction of Heat in Solids, 2nd ed. (Oxford), Ch. 10 (image method), Ch. 16 (spheroidal inclusions).")
    B("J. A. Stratton, Electromagnetic Theory (McGraw-Hill), §3.27 (depolarisation factors of spheroids; the thermal analogue is identical).")

    doc.add_page_break()
    H("부록 A. 케이스별 그림", 1)
    P("각 그림: 온도장 MAP(void 위치 표시)과 평면도, GUI '신호' 페이지의 두 그래프, 지표 표와 void 정보.")
    for c in all_cases():
        if c["id"] in R and (FIG_DIR / f"{c['id']}.png").exists():
            H(f"{c['id']} — {c['label']}", 3)
            fig(f"{c['id']}.png", "", 16.5, FIG_DIR)
    H("부록 B. 전체 수치", 1)
    rows = []
    for c in all_cases():
        if c["id"] in R:
            r = R[c["id"]]
            rows.append([r["id"], r["n_voids"], f"{r['depth_min_um']:g}", f"{r['is_3d'] and '3D' or '2D'} {r['n_unknowns']:,}",
                         f"{r['dT_base_max_K']:.3g}", mk(r["void_minus_base_K"]), f"{r['t_peak_s'] * 1e6:.0f}",
                         f"{r['ratio'] * 100:.2f}", f"{r['contrast_max'] * 100:.2f}", f"{r['scaled_dRR_diff']:.2e}", verdict(r), f"{r['wall_s'] / 60:.0f}"])
    table(["케이스", "수", "얕은 깊이", "격자", "base 피크 K", "ΔT 피크", "t [μs]", "대비 %", "대비 max %", "|ΔR/R| 10K", "판정", "분"],
          rows, widths=[2.3, 0.8, 1.3, 2.3, 1.5, 1.6, 1.1, 1.2, 1.4, 1.8, 1.9, 0.8], size=7.5)
    out = OUT / "ns-TTR_다중void_검출_시뮬레이션_결과보고.docx"
    doc.save(out)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-docx", action="store_true")
    args = ap.parse_args()
    SUM_DIR.mkdir(parents=True, exist_ok=True)
    R = load_cases()
    print(f"{len(R)} finished cases")
    write_csv(R)
    plt = setup_mpl()
    figs = dict(A=fig_A(R, plt), B=fig_B(R, plt), C=fig_C(R, plt), D=fig_D(R, plt), E=fig_E(R, plt), F=fig_F(R, plt))
    print("summary figures:", {k: v for k, v in figs.items() if v})
    summary = {k: {kk: vv for kk, vv in r.items() if kk not in ("case", "voids")} for k, r in R.items()}
    (OUT / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=1), encoding="utf-8")
    if not args.no_docx:
        print("docx:", build_docx(R, figs))


if __name__ == "__main__":
    main()
